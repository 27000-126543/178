import uuid
import logging
from datetime import datetime
from decimal import Decimal
from typing import List, Dict, Optional, Tuple

from sqlalchemy.orm import Session

from models import (
    SessionLocal, TrialBalance, BalanceCheckStatus, Subsidiary,
    ReportSubmission, ReportStatus, OperationLog
)
from config import REPORT_SUBMISSION_DELAY_DAYS

logger = logging.getLogger(__name__)


class TrialBalanceManager:
    """试算平衡表导入与借贷平衡校验"""

    def __init__(self, session: Session = None):
        self.session = session or SessionLocal()

    def import_from_csv(self, file_path: str, company_code: str,
                        period: str) -> Tuple[int, int]:
        import csv
        records = []
        with open(file_path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                row["company_code"] = company_code
                row["period"] = period
                records.append(row)
        return self._save_records(records)

    def import_from_excel(self, file_path: str, company_code: str,
                          period: str) -> Tuple[int, int]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            records = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                record = dict(zip(headers, row))
                record["company_code"] = company_code
                record["period"] = period
                records.append(record)
            wb.close()
            return self._save_records(records)
        except ImportError:
            logger.warning("openpyxl未安装")
            return 0, 0

    def import_from_dict(self, records: List[Dict]) -> Tuple[int, int]:
        return self._save_records(records)

    def _save_records(self, records: List[Dict]) -> Tuple[int, int]:
        saved = 0
        balanced_count = 0
        for rec in records:
            opening = Decimal(str(rec.get("opening_balance", 0) or 0))
            debit = Decimal(str(rec.get("debit_amount", 0) or 0))
            credit = Decimal(str(rec.get("credit_amount", 0) or 0))
            closing = opening + debit - credit
            imbalance = abs(debit - credit) if rec.get("account_type") == "summary" else Decimal("0")

            is_balanced = BalanceCheckStatus.BALANCED
            if imbalance > Decimal("0.01"):
                is_balanced = BalanceCheckStatus.UNBALANCED
            else:
                balanced_count += 1

            tb = TrialBalance(
                id=str(uuid.uuid4()),
                company_code=rec.get("company_code", ""),
                period=rec.get("period", ""),
                account_code=str(rec.get("account_code", "")),
                account_name=str(rec.get("account_name", "")),
                account_type=rec.get("account_type", "detail"),
                opening_balance=opening,
                debit_amount=debit,
                credit_amount=credit,
                closing_balance=closing,
                is_balanced=is_balanced,
                imbalance_amount=imbalance,
            )
            self.session.add(tb)
            saved += 1

        self.session.commit()
        logger.info(f"导入试算平衡表 {saved} 条, 平衡 {balanced_count} 条, 不平衡 {saved - balanced_count} 条")
        return saved, saved - balanced_count

    def validate_balance(self, company_code: str, period: str) -> Dict:
        records = self.session.query(TrialBalance).filter(
            TrialBalance.company_code == company_code,
            TrialBalance.period == period,
        ).all()

        total_debit = sum(r.debit_amount for r in records)
        total_credit = sum(r.credit_amount for r in records)
        imbalance = total_debit - total_credit

        unbalanced = [r for r in records if r.is_balanced == BalanceCheckStatus.UNBALANCED]

        result = {
            "company_code": company_code,
            "period": period,
            "total_debit": float(total_debit),
            "total_credit": float(total_credit),
            "imbalance": float(imbalance),
            "is_balanced": abs(imbalance) <= Decimal("0.01"),
            "total_accounts": len(records),
            "unbalanced_accounts": len(unbalanced),
            "unbalanced_details": [
                {
                    "account_code": r.account_code,
                    "account_name": r.account_name,
                    "debit": float(r.debit_amount),
                    "credit": float(r.credit_amount),
                    "imbalance": float(r.imbalance_amount),
                }
                for r in unbalanced
            ],
        }

        if not result["is_balanced"]:
            self._log_operation("validate_balance", None,
                                f"借贷不平: 借方{total_debit} 贷方{total_credit} 差额{imbalance}",
                                is_anomaly=True)

        return result

    def validate_all_companies(self, period: str) -> List[Dict]:
        subsidiaries = self.session.query(Subsidiary).filter_by(is_active=True).all()
        results = []
        for sub in subsidiaries:
            result = self.validate_balance(sub.code, period)
            results.append(result)
        return results

    def mark_submission(self, company_code: str, period: str,
                        report_type: str, submitted_by: str = "") -> ReportSubmission:
        sub = self.session.query(Subsidiary).filter_by(code=company_code).first()
        deadline = self._get_submission_deadline(period)
        is_late = datetime.utcnow() > deadline

        rs = ReportSubmission(
            id=str(uuid.uuid4()),
            company_code=company_code,
            period=period,
            report_type=report_type,
            status=ReportStatus.SUBMITTED,
            submitted_at=datetime.utcnow(),
            submitted_by=submitted_by,
            is_late=is_late,
        )
        self.session.add(rs)
        self.session.commit()

        if is_late:
            self._log_operation("late_submission", rs.id,
                                f"子公司 {company_code} 报表延迟提交",
                                is_anomaly=True)

        return rs

    def check_late_submissions(self, period: str) -> List[Dict]:
        """检查超2天未提交报表的子公司并催办"""
        deadline = self._get_submission_deadline(period)
        subsidiaries = self.session.query(Subsidiary).filter_by(is_active=True).all()

        late_list = []
        for sub in subsidiaries:
            submission = self.session.query(ReportSubmission).filter(
                ReportSubmission.company_code == sub.code,
                ReportSubmission.period == period,
            ).first()

            if submission and submission.status == ReportStatus.SUBMITTED:
                continue

            if datetime.utcnow() > deadline:
                if submission:
                    submission.is_late = True
                    submission.reminder_count += 1
                    submission.last_reminder_at = datetime.utcnow()
                    self.session.commit()
                else:
                    rs = ReportSubmission(
                        id=str(uuid.uuid4()),
                        company_code=sub.code,
                        period=period,
                        report_type="trial_balance",
                        status=ReportStatus.PENDING,
                        is_late=True,
                        reminder_count=1,
                        last_reminder_at=datetime.utcnow(),
                    )
                    self.session.add(rs)
                    self.session.commit()
                    submission = rs

                late_list.append({
                    "company_code": sub.code,
                    "company_name": sub.name,
                    "cfo_name": sub.cfo_name,
                    "cfo_contact": sub.cfo_contact,
                    "reminder_count": submission.reminder_count,
                    "should_escalate": (
                        datetime.utcnow() - deadline > timedelta(days=REPORT_SUBMISSION_DELAY_DAYS)
                        and submission.reminder_count >= 2
                    ),
                })

        if late_list:
            self._log_operation("check_late_submissions", None,
                                f"发现 {len(late_list)} 家子公司报表延迟",
                                is_anomaly=True)
        return late_list

    def _get_submission_deadline(self, period: str) -> datetime:
        try:
            year, month = int(period[:4]), int(period[5:7])
            from datetime import date
            if month == 12:
                next_month = date(year + 1, 1, 1)
            else:
                next_month = date(year, month + 1, 1)
            from datetime import timedelta
            return datetime.combine(next_month + timedelta(days=REPORT_SUBMISSION_DELAY_DAYS),
                                    datetime.min.time())
        except Exception:
            return datetime.utcnow() + timedelta(days=5)

    def _log_operation(self, op_type: str, target_id: str, detail: str,
                       operator: str = "system", is_anomaly: bool = False):
        log = OperationLog(
            id=str(uuid.uuid4()),
            operation_type=op_type,
            operator=operator,
            target_type="trial_balance",
            target_id=target_id or "",
            detail=detail,
            is_anomaly=is_anomaly,
        )
        self.session.add(log)
        self.session.commit()


from datetime import timedelta
