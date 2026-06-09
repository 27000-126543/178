import uuid
import logging
import os
from datetime import datetime
from decimal import Decimal
from typing import Dict, List, Optional

from sqlalchemy.orm import Session
from sqlalchemy import func, extract, and_

from models import (
    SessionLocal, InternalTransaction, MatchStatus,
    DiscrepancyWorkOrder, WorkOrderStatus, QuarterlyReport, OperationLog
)

logger = logging.getLogger(__name__)

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")


class QuarterlyAnalyzer:
    """季度对账分析报告 - 统计对账完成率、差异率、平均处理时长，对比上季度趋势"""

    def __init__(self, session: Session = None):
        self.session = session or SessionLocal()
        os.makedirs(OUTPUT_DIR, exist_ok=True)

    def generate_quarterly_report(self, year: int, quarter: int) -> Dict:
        month_ranges = {
            1: (1, 3), 2: (4, 6), 3: (7, 9), 4: (10, 12),
        }
        start_month, end_month = month_ranges[quarter]

        total_tx = self.session.query(InternalTransaction).filter(
            extract("year", InternalTransaction.transaction_date) == year,
            extract("month", InternalTransaction.transaction_date).between(start_month, end_month),
        ).count()

        matched_tx = self.session.query(InternalTransaction).filter(
            extract("year", InternalTransaction.transaction_date) == year,
            extract("month", InternalTransaction.transaction_date).between(start_month, end_month),
            InternalTransaction.match_status == MatchStatus.MATCHED,
        ).count()

        discrepancy_count = self.session.query(DiscrepancyWorkOrder).filter(
            extract("year", DiscrepancyWorkOrder.created_at) == year,
            extract("month", DiscrepancyWorkOrder.created_at).between(start_month, end_month),
        ).count()

        resolved_orders = self.session.query(DiscrepancyWorkOrder).filter(
            extract("year", DiscrepancyWorkOrder.created_at) == year,
            extract("month", DiscrepancyWorkOrder.created_at).between(start_month, end_month),
            DiscrepancyWorkOrder.status == WorkOrderStatus.RESOLVED,
            DiscrepancyWorkOrder.resolved_at.isnot(None),
        ).all()

        total_hours = 0
        for wo in resolved_orders:
            if wo.resolved_at and wo.created_at:
                delta = wo.resolved_at - wo.created_at
                total_hours += delta.total_seconds() / 3600

        avg_hours = total_hours / len(resolved_orders) if resolved_orders else 0

        completion_rate = (matched_tx / total_tx * 100) if total_tx > 0 else 100.0
        discrepancy_rate = (discrepancy_count / total_tx * 100) if total_tx > 0 else 0.0

        prev_q, prev_y = self._get_previous_quarter(quarter, year)
        prev_report = self.session.query(QuarterlyReport).filter_by(
            year=prev_y, quarter=prev_q
        ).first()

        result = {
            "year": year,
            "quarter": quarter,
            "total_transactions": total_tx,
            "matched_transactions": matched_tx,
            "discrepancy_count": discrepancy_count,
            "reconciliation_completion_rate": round(completion_rate, 2),
            "discrepancy_rate": round(discrepancy_rate, 2),
            "avg_processing_hours": round(avg_hours, 2),
            "previous_quarter": {
                "year": prev_y,
                "quarter": prev_q,
                "completion_rate": float(prev_report.reconciliation_completion_rate) if prev_report else None,
                "discrepancy_rate": float(prev_report.discrepancy_rate) if prev_report else None,
                "avg_processing_hours": float(prev_report.avg_processing_hours) if prev_report else None,
            },
            "trend": self._calculate_trend(result=None, prev_report=prev_report,
                                            completion_rate=completion_rate,
                                            discrepancy_rate=discrepancy_rate,
                                            avg_hours=avg_hours),
        }

        qr = QuarterlyReport(
            id=str(uuid.uuid4()),
            year=year,
            quarter=quarter,
            reconciliation_completion_rate=Decimal(str(completion_rate)),
            discrepancy_rate=Decimal(str(discrepancy_rate)),
            avg_processing_hours=Decimal(str(avg_hours)),
            total_transactions=total_tx,
            matched_transactions=matched_tx,
            discrepancy_count=discrepancy_count,
        )
        self.session.merge(qr)
        self.session.commit()

        logger.info(f"季度报告生成: {year}Q{quarter} 完成率={completion_rate}%, 差异率={discrepancy_rate}%")
        return result

    def _calculate_trend(self, result, prev_report, completion_rate, discrepancy_rate, avg_hours):
        if not prev_report:
            return {"direction": "initial", "description": "首期报告，无上期对比"}

        prev_completion = float(prev_report.reconciliation_completion_rate)
        prev_discrepancy = float(prev_report.discrepancy_rate)
        prev_avg = float(prev_report.avg_processing_hours)

        completion_change = completion_rate - prev_completion
        discrepancy_change = discrepancy_rate - prev_discrepancy
        avg_change = avg_hours - prev_avg

        direction = "improving"
        if completion_change < -5 or discrepancy_change > 5 or avg_change > 10:
            direction = "declining"
        elif abs(completion_change) <= 5 and abs(discrepancy_change) <= 5:
            direction = "stable"

        return {
            "direction": direction,
            "completion_change": round(completion_change, 2),
            "discrepancy_change": round(discrepancy_change, 2),
            "avg_hours_change": round(avg_change, 2),
            "description": (
                f"完成率{'上升' if completion_change >= 0 else '下降'}{abs(completion_change):.1f}%, "
                f"差异率{'上升' if discrepancy_change >= 0 else '下降'}{abs(discrepancy_change):.1f}%, "
                f"处理时长{'增加' if avg_change >= 0 else '减少'}{abs(avg_change):.1f}小时"
            ),
        }

    def _get_previous_quarter(self, quarter: int, year: int):
        if quarter == 1:
            return 4, year - 1
        return quarter - 1, year

    def generate_trend_chart(self, year: int, quarter: int) -> str:
        """生成季度趋势对比图"""
        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
            plt.rcParams["font.sans-serif"] = ["Arial Unicode MS", "SimHei", "DejaVu Sans"]
            plt.rcParams["axes.unicode_minus"] = False
        except ImportError:
            logger.warning("matplotlib未安装, 跳过趋势图生成")
            return ""

        quarters = []
        completion_rates = []
        discrepancy_rates = []
        avg_hours_list = []

        for q in range(1, 5):
            qr = self.session.query(QuarterlyReport).filter_by(
                year=year, quarter=q
            ).first()
            if qr:
                quarters.append(f"Q{q}")
                completion_rates.append(float(qr.reconciliation_completion_rate))
                discrepancy_rates.append(float(qr.discrepancy_rate))
                avg_hours_list.append(float(qr.avg_processing_hours))
            else:
                quarters.append(f"Q{q}")
                completion_rates.append(0)
                discrepancy_rates.append(0)
                avg_hours_list.append(0)

        fig, axes = plt.subplots(1, 3, figsize=(15, 5))

        axes[0].plot(quarters, completion_rates, "b-o", linewidth=2, markersize=8)
        axes[0].set_title("对账完成率趋势")
        axes[0].set_ylabel("完成率(%)")
        axes[0].set_ylim(0, 105)
        axes[0].grid(True, alpha=0.3)

        axes[1].bar(quarters, discrepancy_rates, color=["#4472C4", "#ED7D31", "#70AD47", "#FFC000"])
        axes[1].set_title("差异率对比")
        axes[1].set_ylabel("差异率(%)")
        axes[1].grid(True, alpha=0.3)

        axes[2].plot(quarters, avg_hours_list, "r-s", linewidth=2, markersize=8)
        axes[2].set_title("平均处理时长趋势")
        axes[2].set_ylabel("小时")
        axes[2].grid(True, alpha=0.3)

        plt.suptitle(f"{year}年度对账分析趋势图", fontsize=14, fontweight="bold")
        plt.tight_layout()

        chart_path = os.path.join(OUTPUT_DIR, f"trend_{year}_Q{quarter}.png")
        plt.savefig(chart_path, dpi=150, bbox_inches="tight")
        plt.close()
        logger.info(f"趋势图生成: {chart_path}")
        return chart_path

    def export_quarterly_report_excel(self, year: int, quarter: int) -> str:
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        except ImportError:
            logger.error("openpyxl未安装")
            return ""

        report = self.generate_quarterly_report(year, quarter)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"季度分析报告_{year}Q{quarter}"

        header_font = Font(name="微软雅黑", size=14, bold=True)
        sub_font = Font(name="微软雅黑", size=11, bold=True)
        data_font = Font(name="微软雅黑", size=10)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        white_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        ws.merge_cells("A1:B1")
        ws["A1"] = f"对账分析报告 - {year}年第{quarter}季度"
        ws["A1"].font = header_font

        data_rows = [
            ("指标", "数值"),
            ("总交易笔数", report["total_transactions"]),
            ("匹配笔数", report["matched_transactions"]),
            ("差异工单数", report["discrepancy_count"]),
            ("对账完成率(%)", report["reconciliation_completion_rate"]),
            ("差异率(%)", report["discrepancy_rate"]),
            ("平均处理时长(小时)", report["avg_processing_hours"]),
        ]

        prev = report.get("previous_quarter", {})
        if prev.get("completion_rate") is not None:
            data_rows.append(("", ""))
            data_rows.append((f"上季度({prev['year']}Q{prev['quarter']})对比", ""))
            data_rows.append(("上季完成率(%)", prev["completion_rate"]))
            data_rows.append(("上季差异率(%)", prev["discrepancy_rate"]))
            data_rows.append(("上季平均时长(小时)", prev["avg_processing_hours"]))

        trend = report.get("trend", {})
        if trend.get("description"):
            data_rows.append(("", ""))
            data_rows.append(("趋势分析", trend["description"]))

        for i, (key, val) in enumerate(data_rows, start=2):
            ws.cell(row=i, column=1, value=key).border = border
            ws.cell(row=i, column=2, value=val).border = border
            if i == 2:
                ws.cell(row=i, column=1).font = white_font
                ws.cell(row=i, column=1).fill = header_fill
                ws.cell(row=i, column=2).font = white_font
                ws.cell(row=i, column=2).fill = header_fill
            else:
                ws.cell(row=i, column=1).font = sub_font if key and not val else data_font
                ws.cell(row=i, column=2).font = data_font

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 25

        chart_path = self.generate_trend_chart(year, quarter)
        if chart_path:
            try:
                img = openpyxl.drawing.image.Image(chart_path)
                img.width = 700
                img.height = 250
                ws.add_image(img, f"A{len(data_rows) + 4}")
            except Exception as e:
                logger.warning(f"Excel嵌入图表失败: {e}")

        filename = f"季度分析报告_{year}Q{quarter}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        filepath = os.path.join(OUTPUT_DIR, filename)
        wb.save(filepath)

        qr = self.session.query(QuarterlyReport).filter_by(year=year, quarter=quarter).first()
        if qr:
            qr.file_path_excel = filepath
            self.session.commit()

        logger.info(f"季度报告Excel导出: {filepath}")
        return filepath
