import uuid
import json
import logging
from datetime import date, datetime, timedelta
from decimal import Decimal
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Dict, Optional, Tuple

from sqlalchemy.orm import Session
from sqlalchemy import and_, or_

from models import (
    SessionLocal, InternalTransaction, MatchStatus,
    DiscrepancyWorkOrder, WorkOrderStatus, EscalationLevel, Subsidiary
)
from config import CONCURRENT_WORKERS, BATCH_SIZE, DIFFERENCE_TOLERANCE

logger = logging.getLogger(__name__)


class TransactionFetcher:
    """从各子公司业务系统抓取内部交易流水"""

    def __init__(self, session: Session = None):
        self.session = session or SessionLocal()

    def fetch_from_api(self, company_code: str, start_date: date, end_date: date) -> List[Dict]:
        logger.info(f"开始抓取子公司 {company_code} 交易流水: {start_date} ~ {end_date}")
        return []

    def fetch_from_csv(self, file_path: str, company_code: str) -> List[Dict]:
        import csv
        transactions = []
        with open(file_path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                row["company_code"] = company_code
                transactions.append(row)
        logger.info(f"从CSV导入 {company_code} 交易 {len(transactions)} 笔")
        return transactions

    def fetch_from_excel(self, file_path: str, company_code: str) -> List[Dict]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            transactions = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                record = dict(zip(headers, row))
                record["company_code"] = company_code
                transactions.append(record)
            wb.close()
            logger.info(f"从Excel导入 {company_code} 交易 {len(transactions)} 笔")
            return transactions
        except ImportError:
            logger.warning("openpyxl未安装，请执行 pip install openpyxl")
            return []

    def save_transactions(self, transactions: List[Dict], batch_id: str = None) -> int:
        if not batch_id:
            batch_id = str(uuid.uuid4())
        saved = 0
        for i in range(0, len(transactions), BATCH_SIZE):
            batch = transactions[i:i + BATCH_SIZE]
            objs = []
            for t in batch:
                obj = InternalTransaction(
                    id=str(uuid.uuid4()),
                    company_code=t.get("company_code", ""),
                    company_name=t.get("company_name", ""),
                    counterparty_code=t.get("counterparty_code", ""),
                    counterparty_name=t.get("counterparty_name", ""),
                    transaction_date=t.get("transaction_date", date.today()),
                    product_code=t.get("product_code", ""),
                    product_name=t.get("product_name", ""),
                    amount=Decimal(str(t.get("amount", 0))),
                    quantity=Decimal(str(t.get("quantity", 0))) if t.get("quantity") else None,
                    direction=t.get("direction", "buy"),
                    match_status=MatchStatus.UNMATCHED,
                    batch_id=batch_id,
                    source_system=t.get("source_system", "manual"),
                    raw_data=json.dumps(t, ensure_ascii=False, default=str),
                )
                objs.append(obj)
            self.session.bulk_save_objects(objs)
            self.session.commit()
            saved += len(objs)
        logger.info(f"保存交易流水 {saved} 笔, batch_id={batch_id}")
        return saved

    def daily_fetch_all(self, target_date: date = None) -> Dict[str, int]:
        if target_date is None:
            target_date = date.today()
        start_date = target_date
        end_date = target_date
        subsidiaries = self.session.query(Subsidiary).filter_by(is_active=True).all()
        results = {}
        for sub in subsidiaries:
            try:
                transactions = self.fetch_from_api(sub.code, start_date, end_date)
                count = self.save_transactions(transactions)
                results[sub.code] = count
            except Exception as e:
                logger.error(f"抓取子公司 {sub.code} 交易失败: {e}")
                results[sub.code] = -1
        return results


class TransactionMatcher:
    """内部交易双向匹配引擎 - 按交易对手、商品编码、金额自动匹配"""

    def __init__(self, session: Session = None):
        self.session = session or SessionLocal()

    def _build_match_key(self, t: InternalTransaction) -> str:
        return f"{t.counterparty_code}|{t.product_code}|{t.amount}"

    def _build_reverse_key(self, t: InternalTransaction) -> str:
        return f"{t.company_code}|{t.product_code}|{t.amount}"

    def match_batch(self, transaction_date: date = None, batch_id: str = None) -> Dict:
        query = self.session.query(InternalTransaction).filter(
            InternalTransaction.match_status == MatchStatus.UNMATCHED
        )
        if transaction_date:
            query = query.filter(InternalTransaction.transaction_date == transaction_date)
        if batch_id:
            query = query.filter(InternalTransaction.batch_id == batch_id)

        unmatched = query.all()
        logger.info(f"待匹配交易 {len(unmatched)} 笔")

        buy_map: Dict[str, List[InternalTransaction]] = {}
        sell_map: Dict[str, List[InternalTransaction]] = {}

        for t in unmatched:
            if t.direction == "buy":
                key = self._build_match_key(t)
                buy_map.setdefault(key, []).append(t)
            else:
                key = self._build_reverse_key(t)
                sell_map.setdefault(key, []).append(t)

        matched_count = 0
        partial_count = 0
        discrepancy_count = 0

        for key, buy_list in buy_map.items():
            if key not in sell_map:
                continue
            sell_list = sell_map[key]

            for buy_tx in buy_list[:]:
                best_match = None
                best_score = 0
                for sell_tx in sell_list[:]:
                    score = self._calculate_match_score(buy_tx, sell_tx)
                    if score > best_score and score >= 0.7:
                        best_score = score
                        best_match = sell_tx

                if best_match:
                    if best_score >= 0.95:
                        buy_tx.match_status = MatchStatus.MATCHED
                        best_match.match_status = MatchStatus.MATCHED
                        buy_tx.matched_transaction_id = best_match.id
                        best_match.matched_transaction_id = buy_tx.id
                        buy_tx.match_score = best_score
                        best_match.match_score = best_score
                        buy_list.remove(buy_tx)
                        sell_list.remove(best_match)
                        matched_count += 1
                    else:
                        buy_tx.match_status = MatchStatus.PARTIAL
                        best_match.match_status = MatchStatus.PARTIAL
                        buy_tx.matched_transaction_id = best_match.id
                        best_match.matched_transaction_id = buy_tx.id
                        buy_tx.match_score = best_score
                        best_match.match_score = best_score
                        self._create_discrepancy(buy_tx, best_match)
                        buy_list.remove(buy_tx)
                        sell_list.remove(best_match)
                        partial_count += 1
                        discrepancy_count += 1

        self.session.commit()
        result = {
            "total_unmatched": len(unmatched),
            "matched": matched_count,
            "partial": partial_count,
            "discrepancy_created": discrepancy_count,
        }
        logger.info(f"匹配完成: {result}")
        return result

    def _calculate_match_score(self, buy_tx: InternalTransaction, sell_tx: InternalTransaction) -> float:
        score = 0.0
        if buy_tx.counterparty_code == sell_tx.company_code:
            score += 0.4
        if buy_tx.product_code == sell_tx.product_code:
            score += 0.3
        amount_diff = abs(buy_tx.amount - sell_tx.amount)
        if buy_tx.amount != 0:
            amount_ratio = float(amount_diff / buy_tx.amount)
            if amount_ratio <= DIFFERENCE_TOLERANCE:
                score += 0.3
            elif amount_ratio <= 0.05:
                score += 0.2
            elif amount_ratio <= 0.1:
                score += 0.1
        if buy_tx.transaction_date and sell_tx.transaction_date:
            date_diff = abs((buy_tx.transaction_date - sell_tx.transaction_date).days)
            if date_diff <= 3:
                score += 0.0
            elif date_diff <= 7:
                score -= 0.1
            else:
                score -= 0.3
        return max(score, 0)

    def _create_discrepancy(self, buy_tx: InternalTransaction, sell_tx: InternalTransaction):
        diff_amount = abs(buy_tx.amount - sell_tx.amount)
        responsible_code = buy_tx.company_code if buy_tx.amount > sell_tx.amount else sell_tx.company_code
        sub = self.session.query(Subsidiary).filter_by(code=responsible_code).first()
        assigned_to = ""
        if sub and sub.finance_staff:
            staff_list = [s.strip() for s in sub.finance_staff.split(",") if s.strip()]
            if staff_list:
                assigned_to = staff_list[0]

        wo = DiscrepancyWorkOrder(
            id=str(uuid.uuid4()),
            transaction_id_buy=buy_tx.id,
            transaction_id_sell=sell_tx.id,
            discrepancy_type="amount_mismatch",
            discrepancy_amount=diff_amount,
            discrepancy_description=(
                f"金额差异: 买方{buy_tx.company_code}金额{buy_tx.amount}, "
                f"卖方{sell_tx.company_code}金额{sell_tx.amount}, "
                f"差异{diff_amount}, 商品{buy_tx.product_code}"
            ),
            responsible_company_code=responsible_code,
            assigned_to=assigned_to,
            status=WorkOrderStatus.OPEN,
        )
        self.session.add(wo)

    def match_concurrent(self, transaction_date: date = None) -> Dict:
        """高并发批量匹配 - 按公司对分组并行处理"""
        query = self.session.query(InternalTransaction).filter(
            InternalTransaction.match_status == MatchStatus.UNMATCHED
        )
        if transaction_date:
            query = query.filter(InternalTransaction.transaction_date == transaction_date)

        unmatched = query.all()
        if not unmatched:
            return {"total_unmatched": 0, "matched": 0, "partial": 0, "discrepancy_created": 0}

        pairs = set()
        for t in unmatched:
            if t.direction == "buy":
                pairs.add((t.company_code, t.counterparty_code))
            else:
                pairs.add((t.counterparty_code, t.company_code))

        total_matched = 0
        total_partial = 0
        total_discrepancy = 0

        with ThreadPoolExecutor(max_workers=CONCURRENT_WORKERS) as executor:
            futures = {}
            for buyer_code, seller_code in pairs:
                future = executor.submit(
                    self._match_company_pair, buyer_code, seller_code, transaction_date
                )
                futures[future] = (buyer_code, seller_code)

            for future in as_completed(futures):
                pair = futures[future]
                try:
                    result = future.result()
                    total_matched += result["matched"]
                    total_partial += result["partial"]
                    total_discrepancy += result["discrepancy_created"]
                except Exception as e:
                    logger.error(f"匹配公司对 {pair} 失败: {e}")

        return {
            "total_unmatched": len(unmatched),
            "matched": total_matched,
            "partial": total_partial,
            "discrepancy_created": total_discrepancy,
        }

    def _match_company_pair(self, buyer_code: str, seller_code: str,
                            transaction_date: date = None) -> Dict:
        session = SessionLocal()
        try:
            query = session.query(InternalTransaction).filter(
                InternalTransaction.match_status == MatchStatus.UNMATCHED
            )
            if transaction_date:
                query = query.filter(InternalTransaction.transaction_date == transaction_date)

            buys = query.filter(
                InternalTransaction.company_code == buyer_code,
                InternalTransaction.counterparty_code == seller_code,
                InternalTransaction.direction == "buy"
            ).all()

            sells = query.filter(
                InternalTransaction.company_code == seller_code,
                InternalTransaction.counterparty_code == buyer_code,
                InternalTransaction.direction == "sell"
            ).all()

            matched = 0
            partial = 0
            discrepancy = 0

            for buy_tx in buys:
                best_match = None
                best_score = 0
                for sell_tx in sells:
                    if sell_tx.match_status != MatchStatus.UNMATCHED:
                        continue
                    score = self._calculate_match_score(buy_tx, sell_tx)
                    if score > best_score and score >= 0.7:
                        best_score = score
                        best_match = sell_tx

                if best_match:
                    if best_score >= 0.95:
                        buy_tx.match_status = MatchStatus.MATCHED
                        best_match.match_status = MatchStatus.MATCHED
                        buy_tx.matched_transaction_id = best_match.id
                        best_match.matched_transaction_id = buy_tx.id
                        buy_tx.match_score = best_score
                        best_match.match_score = best_score
                        sells.remove(best_match)
                        matched += 1
                    else:
                        buy_tx.match_status = MatchStatus.PARTIAL
                        best_match.match_status = MatchStatus.PARTIAL
                        buy_tx.matched_transaction_id = best_match.id
                        best_match.matched_transaction_id = buy_tx.id
                        buy_tx.match_score = best_score
                        best_match.match_score = best_score
                        self._create_discrepancy(buy_tx, best_match)
                        sells.remove(best_match)
                        partial += 1
                        discrepancy += 1

            session.commit()
            return {"matched": matched, "partial": partial, "discrepancy_created": discrepancy}
        except Exception as e:
            session.rollback()
            raise
        finally:
            session.close()
